import csv
from collections import Counter
from io import StringIO
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from flask import current_app
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy import func

from app.domain.client_request import ClientRequest
from app.domain.feedback import Feedback
from app.domain.service import Service
from app.extensions import db
from app.utils.datetime_utils import utc_now


REPORT_TYPES = {
    "requests": "Сводка по заявкам",
    "feedback": "Анализ отзывов",
    "services": "Отчет по услугам",
}
REPORT_FORMATS = {"docx", "xlsx"}

STATUS_LABELS = {
    "new": "Новая",
    "in_progress": "В работе",
    "waiting_client": "Ожидает клиента",
    "completed": "Завершена",
    "cancelled": "Отменена",
}

PRIORITY_LABELS = {
    "low": "Низкий",
    "medium": "Средний",
    "high": "Высокий",
    "urgent": "Срочный",
}

SENTIMENT_LABELS = {
    "positive": "Положительная",
    "neutral": "Нейтральная",
    "negative": "Отрицательная",
}


def get_status_statistics():
    rows = db.session.query(ClientRequest.status, func.count(ClientRequest.id)).group_by(ClientRequest.status).all()
    return {status: count for status, count in rows}


def get_service_statistics():
    rows = (
        db.session.query(Service.title, func.count(ClientRequest.id))
        .outerjoin(ClientRequest, ClientRequest.service_id == Service.id)
        .group_by(Service.id)
        .order_by(func.count(ClientRequest.id).desc(), Service.title.asc())
        .all()
    )
    return rows


def get_sentiment_statistics():
    rows = db.session.query(Feedback.sentiment, func.count(Feedback.id)).group_by(Feedback.sentiment).all()
    return {sentiment: count for sentiment, count in rows}


def get_aspect_statistics():
    rows = db.session.query(Feedback.aspect, func.count(Feedback.id)).group_by(Feedback.aspect).all()
    return {aspect: count for aspect, count in rows}


def generate_requests_csv():
    output = StringIO()
    output.write("\ufeff")
    writer = csv.writer(output, delimiter=";", lineterminator="\n")
    writer.writerow(
        [
            "Номер",
            "Дата создания",
            "Клиент",
            "Компания",
            "Услуга",
            "Название",
            "Статус",
            "Приоритет",
            "Ответственный",
            "Комментариев",
            "Файлов",
            "Тональность последнего отзыва",
            "Аспект последнего отзыва",
        ]
    )

    request_items = ClientRequest.query.order_by(ClientRequest.created_at.desc()).all()
    for request_item in request_items:
        latest_feedback = max(
            request_item.feedback_items,
            key=lambda item: item.created_at,
            default=None,
        )
        writer.writerow(
            [
                request_item.id,
                _format_datetime(request_item.created_at),
                request_item.client.full_name,
                request_item.client.company or "",
                request_item.service.title,
                request_item.title,
                request_item.status_label(),
                request_item.priority_label(),
                request_item.manager.full_name if request_item.manager else "Не назначен",
                len(request_item.comments),
                len(request_item.attachments),
                latest_feedback.sentiment_label() if latest_feedback else "",
                latest_feedback.aspect if latest_feedback else "",
            ]
        )
    return output.getvalue().encode("utf-8")


def generate_report(report_type, report_format):
    if report_type not in REPORT_TYPES:
        raise ValueError("Выбран неизвестный тип отчета.")
    if report_format not in REPORT_FORMATS:
        raise ValueError("Выбран неподдерживаемый формат отчета.")

    report_root = _report_root()
    report_root.mkdir(parents=True, exist_ok=True)
    generated_at = utc_now()
    filename = f"{report_type}_{generated_at.strftime('%Y%m%d_%H%M%S_%f')}.{report_format}"
    report_path = report_root / filename
    report_data = _build_report_data(report_type, generated_at)

    if report_format == "xlsx":
        _write_xlsx(report_path, report_data)
    else:
        _write_docx(report_path, report_data)

    return report_path


def list_generated_reports():
    report_root = _report_root()
    report_root.mkdir(parents=True, exist_ok=True)
    items = []
    for report_path in report_root.iterdir():
        if not report_path.is_file() or report_path.suffix.lower().lstrip(".") not in REPORT_FORMATS:
            continue
        stat = report_path.stat()
        items.append(
            {
                "name": report_path.name,
                "format": report_path.suffix.lower().lstrip(".").upper(),
                "size_kb": max(1, round(stat.st_size / 1024)),
                "created_at": stat.st_mtime,
            }
        )
    return sorted(items, key=lambda item: item["created_at"], reverse=True)


def resolve_generated_report(filename):
    report_root = _report_root()
    target = (report_root / filename).resolve()
    if report_root not in target.parents:
        raise ValueError("Недопустимый путь к отчету.")
    if target.suffix.lower().lstrip(".") not in REPORT_FORMATS:
        raise ValueError("Недопустимый формат отчета.")
    return target


def _report_root():
    return Path(current_app.config["REPORT_FOLDER"]).resolve()


def _build_report_data(report_type, generated_at):
    if report_type == "requests":
        sections = _request_report_sections()
    elif report_type == "feedback":
        sections = _feedback_report_sections()
    else:
        sections = _service_report_sections()
    return {
        "title": REPORT_TYPES[report_type],
        "generated_at": generated_at,
        "sections": sections,
    }


def _request_report_sections():
    request_items = ClientRequest.query.order_by(ClientRequest.created_at.desc()).all()
    status_counts = Counter(item.status for item in request_items)
    priority_counts = Counter(item.priority for item in request_items)
    manager_counts = Counter(item.manager.full_name if item.manager else "Не назначен" for item in request_items)
    completed_durations = [
        (item.updated_at - item.created_at).total_seconds() / 3600
        for item in request_items
        if item.status == "completed" and item.updated_at and item.created_at
    ]
    average_hours = round(sum(completed_durations) / len(completed_durations), 2) if completed_durations else "Нет данных"

    return [
        {
            "title": "Общие показатели",
            "headers": ["Показатель", "Значение"],
            "rows": [
                ["Всего заявок", len(request_items)],
                ["Среднее время завершения, часов", average_hours],
            ],
        },
        {
            "title": "Распределение по статусам",
            "headers": ["Статус", "Количество"],
            "rows": [[label, status_counts.get(code, 0)] for code, label in STATUS_LABELS.items()],
        },
        {
            "title": "Распределение по приоритетам",
            "headers": ["Приоритет", "Количество"],
            "rows": [[label, priority_counts.get(code, 0)] for code, label in PRIORITY_LABELS.items()],
        },
        {
            "title": "Распределение по ответственным",
            "headers": ["Ответственный", "Количество"],
            "rows": [[manager, count] for manager, count in sorted(manager_counts.items())],
        },
        {
            "title": "Реестр заявок",
            "headers": [
                "Номер",
                "Дата",
                "Клиент",
                "Услуга",
                "Название",
                "Статус",
                "Приоритет",
                "Ответственный",
            ],
            "rows": [
                [
                    item.id,
                    _format_datetime(item.created_at),
                    item.client.full_name,
                    item.service.title,
                    item.title,
                    item.status_label(),
                    item.priority_label(),
                    item.manager.full_name if item.manager else "Не назначен",
                ]
                for item in request_items
            ],
        },
    ]


def _feedback_report_sections():
    feedback_items = Feedback.query.order_by(Feedback.created_at.desc()).all()
    sentiment_counts = Counter(item.sentiment for item in feedback_items)
    aspect_counts = Counter(item.aspect for item in feedback_items)
    daily_counts = Counter(item.created_at.strftime("%d.%m.%Y") for item in feedback_items if item.created_at)

    return [
        {
            "title": "Распределение тональности",
            "headers": ["Тональность", "Количество"],
            "rows": [[label, sentiment_counts.get(code, 0)] for code, label in SENTIMENT_LABELS.items()],
        },
        {
            "title": "Тематические аспекты",
            "headers": ["Аспект", "Количество"],
            "rows": [[aspect, count] for aspect, count in sorted(aspect_counts.items())],
        },
        {
            "title": "Динамика отзывов по датам",
            "headers": ["Дата", "Количество"],
            "rows": [[date_value, count] for date_value, count in sorted(daily_counts.items())],
        },
        {
            "title": "Реестр отзывов",
            "headers": ["Дата", "Заявка", "Клиент", "Тональность", "Аспект", "Текст"],
            "rows": [
                [
                    _format_datetime(item.created_at),
                    item.request_id,
                    item.author.full_name,
                    item.sentiment_label(),
                    item.aspect,
                    item.text,
                ]
                for item in feedback_items
            ],
        },
    ]


def _service_report_sections():
    service_items = Service.query.order_by(Service.title.asc()).all()
    rows = []
    for service in service_items:
        request_items = list(service.requests)
        feedback_items = [feedback for item in request_items for feedback in item.feedback_items]
        completed_durations = [
            (item.updated_at - item.created_at).total_seconds() / 3600
            for item in request_items
            if item.status == "completed" and item.updated_at and item.created_at
        ]
        average_hours = round(sum(completed_durations) / len(completed_durations), 2) if completed_durations else "Нет данных"
        positive_share = (
            round(sum(item.sentiment == "positive" for item in feedback_items) * 100 / len(feedback_items), 1)
            if feedback_items
            else "Нет данных"
        )
        rows.append(
            [
                service.title,
                len(request_items),
                len(feedback_items),
                positive_share,
                average_hours,
            ]
        )

    return [
        {
            "title": "Популярность услуг",
            "headers": [
                "Услуга",
                "Заявок",
                "Отзывов",
                "Положительных отзывов, %",
                "Среднее время завершения, часов",
            ],
            "rows": sorted(rows, key=lambda row: (-row[1], row[0])),
        }
    ]


def _write_xlsx(report_path, report_data):
    workbook = Workbook()
    workbook.remove(workbook.active)

    for index, section in enumerate(report_data["sections"], start=1):
        worksheet = workbook.create_sheet(title=_sheet_title(section["title"], index))
        worksheet["A1"] = report_data["title"]
        worksheet["A1"].font = Font(size=16, bold=True, color="1F3A5F")
        worksheet["A2"] = f"Сформирован: {_format_datetime(report_data['generated_at'])}"
        worksheet["A2"].font = Font(italic=True, color="5B6573")
        worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(section["headers"]))
        worksheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(section["headers"]))
        worksheet.sheet_view.showGridLines = False
        worksheet.page_setup.orientation = "landscape"
        worksheet.page_setup.fitToWidth = 1
        worksheet.page_setup.fitToHeight = 0
        worksheet.sheet_properties.pageSetUpPr.fitToPage = True
        worksheet.print_title_rows = "1:4"
        worksheet.append([])
        worksheet.append(section["headers"])

        header_row = 4
        for cell in worksheet[header_row]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill("solid", fgColor="1F4E78")
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)

        for row in section["rows"]:
            worksheet.append([_excel_value(value) for value in row])

        for row in worksheet.iter_rows(min_row=header_row + 1):
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        worksheet.freeze_panes = "A5"
        worksheet.auto_filter.ref = f"A{header_row}:{get_column_letter(len(section['headers']))}{max(header_row, worksheet.max_row)}"
        _fit_worksheet_columns(worksheet)

    workbook.save(report_path)


def _write_docx(report_path, report_data):
    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    section.top_margin = Inches(0.55)
    section.bottom_margin = Inches(0.55)
    section.left_margin = Inches(0.6)
    section.right_margin = Inches(0.6)

    normal_style = document.styles["Normal"]
    normal_style.font.name = "Arial"
    normal_style.font.size = Pt(10)
    heading_style = document.styles["Heading 1"]
    heading_style.font.name = "Arial"
    heading_style.font.color.rgb = RGBColor(0, 0, 0)
    heading_style.font.size = Pt(14)

    title = document.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title.add_run(report_data["title"])
    title_run.bold = True
    title_run.font.name = "Arial"
    title_run.font.size = Pt(20)
    title_run.font.color.rgb = RGBColor(0, 0, 0)
    generated_paragraph = document.add_paragraph(
        f"Дата формирования: {_format_datetime(report_data['generated_at'])}"
    )
    generated_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for section in report_data["sections"]:
        if section["title"].startswith("Реестр"):
            document.add_page_break()
        section_heading = document.add_heading(section["title"], level=1)
        section_heading.paragraph_format.keep_with_next = True
        table = document.add_table(rows=1, cols=len(section["headers"]))
        table.style = "Table Grid"
        table.autofit = False
        column_widths = _docx_column_widths(section["headers"])
        for index, header in enumerate(section["headers"]):
            cell = table.rows[0].cells[index]
            cell.width = Inches(column_widths[index])
            cell.text = str(header)
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.name = "Arial"
                run.font.size = Pt(8)
        _repeat_table_header(table.rows[0])
        _prevent_row_split(table.rows[0])

        rows = section["rows"] or [["Нет данных"] + [""] * (len(section["headers"]) - 1)]
        for row in rows:
            table_row = table.add_row()
            _prevent_row_split(table_row)
            cells = table_row.cells
            for index, value in enumerate(row):
                cells[index].width = Inches(column_widths[index])
                cells[index].text = str(value if value is not None else "")
                for run in cells[index].paragraphs[0].runs:
                    run.font.name = "Arial"
                    run.font.size = Pt(8)

        document.add_paragraph()

    document.save(report_path)


def _sheet_title(title, index):
    safe_title = title.replace("/", "-").replace("\\", "-").replace("?", "")
    return f"{index}. {safe_title}"[:31]


def _fit_worksheet_columns(worksheet):
    for column_cells in worksheet.columns:
        max_length = max(len(str(cell.value or "")) for cell in column_cells)
        column_letter = get_column_letter(column_cells[0].column)
        worksheet.column_dimensions[column_letter].width = min(max(max_length + 2, 12), 48)


def _docx_column_widths(headers):
    preferred_weights = {
        "Номер": 0.7,
        "Дата": 1.1,
        "Заявка": 0.8,
        "Клиент": 1.5,
        "Услуга": 1.6,
        "Название": 2.2,
        "Статус": 1.0,
        "Приоритет": 1.0,
        "Ответственный": 1.6,
        "Тональность": 1.25,
        "Аспект": 1.3,
        "Текст": 3.6,
        "Показатель": 3.0,
        "Значение": 1.4,
    }
    weights = [preferred_weights.get(str(header), max(1.0, min(len(str(header)) / 9, 2.5))) for header in headers]
    total_weight = sum(weights)
    available_width = 10.0
    return [available_width * weight / total_weight for weight in weights]


def _repeat_table_header(row):
    table_properties = row._tr.get_or_add_trPr()
    repeat_element = OxmlElement("w:tblHeader")
    repeat_element.set(qn("w:val"), "true")
    table_properties.append(repeat_element)


def _prevent_row_split(row):
    table_properties = row._tr.get_or_add_trPr()
    no_split_element = OxmlElement("w:cantSplit")
    table_properties.append(no_split_element)


def _format_datetime(value):
    return value.strftime("%d.%m.%Y %H:%M") if value else ""


def _excel_value(value):
    if value is None:
        return ""
    return value
