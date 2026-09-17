from io import BytesIO

from docx import Document
from openpyxl import load_workbook
from sqlalchemy import inspect

from app.domain.attachment import Attachment
from app.extensions import db


def test_required_database_tables_exist(app):
    with app.app_context():
        tables = set(inspect(db.engine).get_table_names())
    required = {
        "users",
        "services",
        "client_requests",
        "request_comments",
        "feedback",
        "audit_log",
        "projects",
        "attachments",
        "notifications",
        "system_settings",
        "contact_messages",
    }
    assert required <= tables


def test_attachment_upload_and_download(app, client, login):
    login("client", "client1234")
    response = client.post(
        "/client/requests/1/attachments",
        data={"attachment": (BytesIO("Акт осмотра".encode()), "inspection.txt")},
        content_type="multipart/form-data",
        follow_redirects=True,
    )
    assert response.status_code == 200
    assert "Файл добавлен к заявке" in response.get_data(as_text=True)

    with app.app_context():
        attachment = Attachment.query.filter_by(original_name="inspection.txt").one()
        attachment_id = attachment.id

    download = client.get(f"/client/attachments/{attachment_id}/download")
    assert download.status_code == 200
    assert download.data == "Акт осмотра".encode()


def test_manager_can_generate_csv_report(client, login):
    login("manager", "manager1234")
    response = client.get("/manager/reports/export.csv")
    assert response.status_code == 200
    assert "attachment" in response.headers["Content-Disposition"]
    report = response.data.decode("utf-8-sig")
    assert "Номер;Дата создания;Клиент" in report
    assert "Проверка качества работ на объекте" in report


def test_manager_can_generate_xlsx_reports(app, client, login):
    login("manager", "manager1234")
    expected_titles = {
        "requests": "Сводка по заявкам",
        "feedback": "Анализ отзывов",
        "services": "Отчет по услугам",
    }

    for report_type, expected_title in expected_titles.items():
        response = client.post(
            "/manager/reports/generate",
            data={"report_type": report_type, "report_format": "xlsx"},
        )
        assert response.status_code == 200
        assert response.data.startswith(b"PK")
        assert f"{report_type}_" in response.headers["Content-Disposition"]
        assert ".xlsx" in response.headers["Content-Disposition"]

        workbook = load_workbook(BytesIO(response.data), read_only=True)
        assert workbook.sheetnames
        assert workbook[workbook.sheetnames[0]]["A1"].value == expected_title

    report_files = sorted(app.config["REPORT_FOLDER"].glob("*.xlsx"))
    assert len(report_files) == len(expected_titles)


def test_manager_can_generate_and_download_docx_reports(app, client, login):
    login("manager", "manager1234")
    expected_titles = {
        "requests": "Сводка по заявкам",
        "feedback": "Анализ отзывов",
        "services": "Отчет по услугам",
    }

    for report_type, expected_title in expected_titles.items():
        response = client.post(
            "/manager/reports/generate",
            data={"report_type": report_type, "report_format": "docx"},
        )
        assert response.status_code == 200
        assert response.data.startswith(b"PK")
        assert f"{report_type}_" in response.headers["Content-Disposition"]
        assert ".docx" in response.headers["Content-Disposition"]

        document = Document(BytesIO(response.data))
        assert document.paragraphs[0].text == expected_title
        assert document.tables

    report_files = sorted(app.config["REPORT_FOLDER"].glob("*.docx"))
    assert len(report_files) == len(expected_titles)
    selected_report = report_files[0]
    download = client.get(f"/manager/reports/files/{selected_report.name}")
    assert download.status_code == 200
    assert download.data == selected_report.read_bytes()

    page = client.get("/manager/reports")
    assert page.status_code == 200
    assert selected_report.name in page.get_data(as_text=True)


def test_client_cannot_generate_office_report(client, login):
    login("client", "client1234")
    response = client.post(
        "/manager/reports/generate",
        data={"report_type": "requests", "report_format": "xlsx"},
    )
    assert response.status_code == 403
